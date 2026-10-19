"""The LDS Sexuality Survey (June 2025) → aggregate priors, added to ucsim/priors/survey_priors.json in place.

    python -m ucsim.ingest.sexuality            # needs data/raw/lds-sexuality_2025-06-29.xlsx (never committed)

Incremental on purpose: the other four surveys' sections were built from JSON exports that the CSVs in data/raw/
don't reproduce exactly, so this module rewrites only its own keys (blocks.lds_sexuality, targets.lds_sex_*,
tables.lds_sex_*, meta.sources["lds-sexuality"]) and leaves everything else byte-for-byte as it was.

Privacy: the same boundary as ucsim.ingest.priors — only aggregates, every stored cell has ≥ K_MIN respondents,
no free text (those columns are never read; see survey_spec.SEXUALITY_SPEC), no respondent rows. The extra sheets
(Fantasy, Favorites, Location, Dropdowns) are never opened.
"""
from __future__ import annotations

import json
import re
import sys
import warnings
from collections import Counter
from datetime import date

import numpy as np
import polars as pl

from .priors import K_MIN, PRIORS_DIR, _nband, _with_bands, copula_block, summarize
from .survey_spec import NOWFREQ, REL_MAP, SEXUALITY_SPEC, STAGE5, WHEN_TIERS
from .surveys import RAW_DIR, _code_value, _norm

SLUG = "lds-sexuality"
BLOCK = "lds_sexuality"
PREFIX = "lds_sex_"
AGE_SCOPE = (18, 35)          # the model's population
TIERS = list(WHEN_TIERS)      # touch, underclothes, naked, sex


def raw_file():
    p = RAW_DIR / SEXUALITY_SPEC["file"]
    return p if p.exists() else None


def _key(h) -> str:
    return re.sub(r"\s+", " ", _norm(h or "")).strip()


def _columns(header: list) -> dict[str, int]:
    """spec key -> column index: an exact header match, else the one header the key is a prefix of."""
    hs = [_key(h) for h in header]
    out = {}
    for k in SEXUALITY_SPEC["vars"]:
        kk = _key(k)
        exact = [j for j, h in enumerate(hs) if h == kk]
        pref = [j for j, h in enumerate(hs) if h.startswith(kk)]
        hit = exact or pref
        if len(hit) != 1:
            raise ValueError(f"survey column {k!r}: {len(hit)} matches in the header")
        out[k] = hit[0]
    return out


def load() -> pl.DataFrame:
    """Respondent-level, harmonized, in memory only (never written anywhere)."""
    path = raw_file()
    if path is None:
        raise FileNotFoundError(f"The LDS Sexuality Survey export isn't in {RAW_DIR} ({SEXUALITY_SPEC['file']}); "
                                "the owner supplies it — it is never committed.")
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]                      # the responses; the other sheets are free text, never read
    it = ws.iter_rows(values_only=True)
    cols = _columns(list(next(it)))
    rows = []
    for r in it:
        row: dict = {}
        for k, (name, vtype, opts) in SEXUALITY_SPEC["vars"].items():
            raw = r[cols[k]]
            if isinstance(raw, float) and vtype != "num" and raw.is_integer():
                raw = int(raw)
            val = _code_value(vtype, opts, raw)
            if vtype == "multi":
                for flag in opts["flags"].values():
                    row[flag] = None if not isinstance(val, dict) else val[flag]
            else:
                row[name] = val
        rows.append(row)
    wb.close()
    return derive(pl.DataFrame(rows, infer_schema_length=None))


def _ladder(df: pl.DataFrame, tier: str) -> pl.Series:
    """The stage a tier becomes acceptable: the median (rounded down) of its items, when at least half are answered."""
    items = [f"when_{tier}_{j}" for j in range(len(WHEN_TIERS[tier]))]
    a = np.column_stack([df[c].cast(pl.Float64).to_numpy() for c in items])
    ok = np.isfinite(a).sum(1) * 2 >= a.shape[1]
    with np.errstate(all="ignore"), warnings.catch_warnings():
        warnings.simplefilter("ignore", RuntimeWarning)
        med = np.floor(np.nanmedian(a, axis=1))
    med = np.minimum(med, 3)                       # "never" (≈0.3%) pools with "marriage"
    return pl.Series(f"ok_{tier}", np.where(ok, med, np.nan)).fill_nan(None)


def _now_freq(df: pl.DataFrame, now_flags: list[str], freqs: list[str], periods_answered: str) -> pl.Series:
    """0 none / 1 every few months or less / 2 monthly / 3 weekly+ (the model's porn_use scale), from 'Now (in the
    last year)' and the current-frequency item; missing when the period item wasn't answered."""
    now = np.zeros(df.height, bool)
    for f in now_flags:
        now |= df[f].fill_null(0).to_numpy() == 1
    fr = np.full(df.height, -1.0)
    for f in freqs:
        fr = np.maximum(fr, df[f].cast(pl.Float64).fill_null(-1).to_numpy())
    level = np.where(fr >= 2, 3, np.where(fr >= 1, 2, 1))
    out = np.where(now, level, 0).astype(float)
    out[df[periods_answered].is_null().to_numpy()] = np.nan
    return pl.Series(out).fill_nan(None)


def derive(df: pl.DataFrame) -> pl.DataFrame:
    sex = df["sex_raw"].map_elements(lambda s: {"male": "M", "female": "F"}.get(_norm(s)), return_dtype=pl.Utf8)
    rel = df["rel_raw"].map_elements(lambda s: REL_MAP.get(_norm(s)), return_dtype=pl.Utf8)
    df = df.with_columns(sex=sex, rel=rel,
                         lds=df["membership"].map_elements(lambda s: int(_norm(s) == "current member"), return_dtype=pl.Int64))
    orient = df["orientation"].map_elements(
        lambda s: {"straight": "straight", "bisexual": "bi", "pansexual": "bi", "gay": "gay", "asexual": "asexual"}.get(_norm(s), "other"),
        return_dtype=pl.Utf8)
    rel_level = pl.col("rel").replace_strict({"single": 0, "divorced": 0, "dating": 1, "engaged": 2, "married": 3},
                                             default=None, return_dtype=pl.Int64)
    df = df.with_columns(orientation=orient, rel_level=rel_level,
                         activity=pl.when(pl.col("lds") == 1).then(pl.col("activity")).otherwise(None),
                         married=(pl.col("rel") == "married").cast(pl.Int64))
    df = df.with_columns(*[_ladder(df, t) for t in TIERS])
    late = df["porn_wrong"].is_not_null()          # reached the attitude page (the survey loses ~27% before it)
    bishop = (df["bishop_porn"].fill_null(0) + df["bishop_mast"].fill_null(0)) > 0
    df = df.with_columns(
        porn_use=_now_freq(df, ["porn_img_now", "porn_vid_now"], ["porn_img_freq", "porn_vid_freq"], "porn_img_now"),
        masturbation=_now_freq(df, ["mast_now"], ["mast_freq"], "mast_now"),
        porn_ever=pl.when(pl.col("porn_img_never").is_null()).then(None).otherwise(
            ((pl.col("porn_img_teen") + pl.col("porn_img_mission") + pl.col("porn_img_recent") + pl.col("porn_img_now")
              + pl.col("porn_vid_teen") + pl.col("porn_vid_mission") + pl.col("porn_vid_recent") + pl.col("porn_vid_now")) > 0).cast(pl.Int64)),
        bishop_any=pl.Series(np.where(late.to_numpy(), bishop.to_numpy().astype(float), np.nan)).fill_nan(None),
        bishop_porn_any=pl.Series(np.where(late.to_numpy(), df["bishop_porn"].fill_null(0).to_numpy().astype(float), np.nan)).fill_nan(None),
        bishop_mast_any=pl.Series(np.where(late.to_numpy(), df["bishop_mast"].fill_null(0).to_numpy().astype(float), np.nan)).fill_nan(None),
        premarital_any=(pl.col("premarital_partners") > 0).cast(pl.Int64),
    )
    df = df.with_columns(porn_weekly=(pl.col("porn_use") >= 3).cast(pl.Int64),
                         porn_now=(pl.col("porn_use") >= 1).cast(pl.Int64),
                         **{f"{d}_agree": (pl.col(d) >= 3).cast(pl.Int64) for d in DEALBREAKERS})
    return df


DEALBREAKERS = ["db_watches_porn", "db_past_porn", "db_masturbates", "db_not_virgin", "db_asexual"]
BLOCK_VARS = {
    # bridges (shared with the ProvoStats block the population is drawn from first)
    "age": "num", "rel_level": "ord", "activity": "ord", "mission": "bin", "n_kissed": "num", "n_ncmo": "num",
    "n_relationships": "num", "premarital_partners": "num",
    # this survey's own
    "sex_drive": "ord", **{f"ok_{t}": "ord" for t in TIERS}, "sex_ed_home": "ord", "sexual_shame": "ord",
    "porn_wrong": "ord", "porn_use": "ord", "masturbation": "ord", "bishop_any": "bin",
    **{d: "ord" for d in DEALBREAKERS}, "talk_discomfort": "ord",
    "sex_satisfaction": "ord", "marriage_satisfaction": "ord", "desire_gap": "ord",
}


def scope(df: pl.DataFrame) -> pl.DataFrame:
    return df.filter((pl.col("age") >= AGE_SCOPE[0]) & (pl.col("age") <= AGE_SCOPE[1]) & pl.col("sex").is_in(["M", "F"]))


def _shares(df: pl.DataFrame, col: str) -> dict:
    """Make-up of the sample: shares of levels with ≥ K_MIN respondents (rarer levels pooled into 'other')."""
    vc = Counter(x for x in df[col].to_list() if x is not None)
    tot = sum(vc.values())
    out = {str(k): round(v / tot, 3) for k, v in vc.items() if v >= K_MIN}
    rest = tot - sum(v for v in vc.values() if v >= K_MIN)
    if rest >= K_MIN:
        out["other"] = round(rest / tot, 3)
    return out


def _stage_table(df: pl.DataFrame, by: str | None) -> dict:
    """Share of people who think each tier is acceptable at each stage (stage labels as words), by a group."""
    words = ["on a date", "committed relationship", "engaged", "married", "never"]
    groups = [("all", df)] if by is None else [(str(k[0] if isinstance(k, tuple) else k), s) for k, s in df.group_by(by)]
    out = {}
    for g, sub in groups:
        if g == "None":
            continue
        row = {}
        for t in TIERS:
            x = sub[f"ok_{t}"].drop_nulls().to_numpy().astype(int)
            if len(x) < 30:
                continue
            cnt = np.bincount(x, minlength=5)
            row[t] = {w: round(float(c / len(x)), 3) for w, c in zip(words, cnt) if c >= K_MIN}
        if row:
            out[g] = row
    return out


def pool_rare_levels(df: pl.DataFrame, cols: list[str]) -> pl.DataFrame:
    """Per sex, recode answer levels held by fewer than K_MIN people to the nearest level that has enough, so no
    stored marginal (even a smoothed one) stands for a tiny group."""
    out = []
    for s, sub in df.group_by("sex"):
        for c in cols:
            x = sub[c].cast(pl.Float64).to_numpy()
            ok = np.isfinite(x)
            levels, cnt = np.unique(x[ok].astype(int), return_counts=True)
            big = levels[cnt >= K_MIN]
            if len(big) and len(big) < len(levels):
                xi = x[ok].astype(int)
                x[ok] = big[np.abs(xi[:, None] - big[None, :]).argmin(1)]
                sub = sub.with_columns(pl.Series(c, x).fill_nan(None))
        out.append(sub)
    return pl.concat(out, how="diagonal_relaxed")


def build(write: bool = True) -> dict:
    df = _with_bands(scope(load()))
    df = pool_rare_levels(df, [v for v, k in BLOCK_VARS.items() if k in ("ord", "bin") and v not in ("rel_level", "activity", "mission")])
    n = df.height
    unmarried = df.filter(pl.col("married") == 0)
    married = df.filter(pl.col("married") == 1)
    block = copula_block(df, BLOCK, BLOCK_VARS,
                         "The LDS Sexuality Survey (June 2025): drive, the stage each tier of physical intimacy becomes "
                         "acceptable, sex education at home, shame, porn and masturbation now, talking to a bishop, the "
                         "five dating dealbreakers, and married couples' satisfaction and desire gap.")
    T = {
        "n_kissed_by_sex_activity": summarize(df, "n_kissed", ["sex", "activity"]),
        "n_kissed_by_sex_age": summarize(df, "n_kissed", ["sex", "age_band"]),
        "n_ncmo_by_sex_activity": summarize(df, "n_ncmo", ["sex", "activity"]),
        "n_relationships_by_sex_age": summarize(df, "n_relationships", ["sex", "age_band"]),
        "age_first_kiss_by_sex_activity": summarize(df, "age_first_kiss", ["sex", "activity"]),
        "age_first_relationship_by_sex_activity": summarize(df, "age_first_relationship", ["sex", "activity"]),
        "premarital_sex_unmarried_by_activity": summarize(unmarried, "premarital_any", ["activity"]),
        "sex_with_spouse_before_marriage_by_activity": summarize(married, "sex_with_spouse_before", ["activity"]),
        "porn_now_by_sex_activity": summarize(df, "porn_now", ["sex", "activity"]),
        "porn_weekly_by_sex_activity": summarize(df, "porn_weekly", ["sex", "activity"]),
        "porn_ever_by_sex": summarize(df, "porn_ever", ["sex"]),
        "bishop_porn_by_sex": summarize(df, "bishop_porn_any", ["sex"]),
        "bishop_mast_by_sex": summarize(df, "bishop_mast_any", ["sex"]),
        "sex_satisfaction_married_by_sex": summarize(married, "sex_satisfaction", ["sex"]),
        "marriage_satisfaction_married_by_sex": summarize(married, "marriage_satisfaction", ["sex"]),
        "dealbreakers_by_sex": {d: summarize(df, f"{d}_agree", ["sex"]) for d in DEALBREAKERS},
    }
    tables = {"ladder_by_sex": _stage_table(df, "sex"), "ladder_by_activity": _stage_table(df, "activity"),
              "ladder_all": _stage_table(df, None)}
    source = {"title": SEXUALITY_SPEC["title"], "n_band": _nband(n), "fielded": "2025-06",
              "scope": f"ages {AGE_SCOPE[0]}-{AGE_SCOPE[1]}, men and women",
              "make_up": {"sex": _shares(df, "sex"), "rel": _shares(df, "rel"), "activity": _shares(df, "activity"),
                          "utah": _shares(df.with_columns(pl.col("origin_utah").cast(pl.Utf8)), "origin_utah"),
                          "age_band": _shares(df, "age_band")}}
    out = {"block": block, "targets": T, "tables": tables, "source": source, "n_in_scope": n}
    if write:
        merge_into_priors(out)
    return out


def merge_into_priors(out: dict, path=None) -> dict:
    """Replace only this survey's sections; everything else in the file stays as it was."""
    path = path or PRIORS_DIR / "survey_priors.json"
    pri = json.loads(path.read_text(encoding="utf-8"))
    pri["blocks"][BLOCK] = out["block"]
    for k in [k for k in pri["targets"] if k.startswith(PREFIX)]:
        del pri["targets"][k]
    for k in [k for k in pri["tables"] if k.startswith(PREFIX)]:
        del pri["tables"][k]
    pri["targets"].update({PREFIX + k: v for k, v in out["targets"].items()})
    pri["tables"].update({PREFIX + k: v for k, v in out["tables"].items()})
    pri["meta"]["sources"][SLUG] = {**out["source"], "built": date.today().isoformat()}
    path.write_text(json.dumps(pri, indent=1), encoding="utf-8")
    return pri


if __name__ == "__main__":
    if raw_file() is None:
        print(f"Missing {RAW_DIR / SEXUALITY_SPEC['file']} (the owner supplies it; it is never committed).")
        sys.exit(1)
    o = build()
    s = o["source"]
    print(f"in scope: {o['n_in_scope']} respondents ({s['scope']}); make-up {json.dumps(s['make_up'])}")
    print(f"block {BLOCK}: {len(o['block']['variables'])} variables, min pair n {o['block']['min_pair_n']}")
